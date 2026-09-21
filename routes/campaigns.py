import json
import re
import csv
import io
import openpyxl
import os
from flask import Blueprint, render_template, request, jsonify, redirect, url_for, flash, Response
from flask_login import login_required, current_user

from models import db, Campaign, CampaignRecipient
from helpers.decorators import subscription_required
from helpers.quota import _daily_quota_status, _create_quota_overflow_schedule, DAILY_SEND_LIMIT, _utc_tomorrow_start
from helpers.emails import extract_emails, derive_name_from_email, _safe_fetch, resolve_recipient_name

campaigns_bp = Blueprint("campaigns", __name__)

@campaigns_bp.route("/campaign/new")
@login_required
@subscription_required
def campaign_new():
    return render_template(
        "campaign_new.html",
        daily_quota=_daily_quota_status(current_user.id),
    )

@campaigns_bp.route("/campaign/<int:campaign_id>")
@login_required
@subscription_required
def campaign_detail(campaign_id):
    campaign = Campaign.query.filter_by(id=campaign_id, user_id=current_user.id).first_or_404()
    recipients = CampaignRecipient.query.filter_by(campaign_id=campaign.id).all()
    if recipients:
        sent = sum(1 for recipient in recipients if recipient.status == "sent")
        failed = sum(1 for recipient in recipients if recipient.status == "failed")
        pending = sum(1 for recipient in recipients if recipient.status == "pending")
    else:
        sent = campaign.sent_ok or 0
        failed = campaign.sent_fail or 0
        pending = max((campaign.total or 0) - sent - failed, 0)
    attempted = sent + failed
    metrics = {
        "total": max(campaign.total or 0, sent + failed + pending),
        "sent": sent,
        "failed": failed,
        "pending": pending,
        "sent_rate": round(sent / attempted * 100, 1) if attempted else None,
    }
    return render_template(
        "campaign_detail.html",
        campaign=campaign,
        recipients=recipients,
        metrics=metrics,
    )

@campaigns_bp.route("/campaign/<int:campaign_id>/resend-failed")
@login_required
@subscription_required
def campaign_resend_failed(campaign_id):
    campaign = Campaign.query.filter_by(id=campaign_id, user_id=current_user.id).first_or_404()
    failed = [
        r.email
        for r in CampaignRecipient.query.filter_by(campaign_id=campaign.id, status="failed").all()
    ]
    if not failed:
        flash("This campaign has no failed recipients to resend to.", "error")
        return redirect(url_for("campaigns.campaign_detail", campaign_id=campaign.id))
    prefill = {
        "name": f"{campaign.name} (retry)",
        "subject": campaign.subject or "",
        "body": campaign.body or "",
        "emails": failed,
    }
    return render_template(
        "campaign_new.html",
        prefill=prefill,
        daily_quota=_daily_quota_status(current_user.id),
    )

@campaigns_bp.route("/campaign/<int:campaign_id>/retry", methods=["POST"])
@login_required
@subscription_required
def campaign_retry(campaign_id):
    from helpers.scheduler_tasks import _send_campaign_background
    from extensions import _scheduler

    campaign = Campaign.query.filter_by(id=campaign_id, user_id=current_user.id).first_or_404()

    if campaign.status in ("queued", "sending"):
        flash("This campaign is already sending. Please wait for it to finish.", "error")
        return redirect(url_for("campaigns.campaign_detail", campaign_id=campaign.id))

    pending = CampaignRecipient.query.filter_by(campaign_id=campaign_id, status="pending").all()
    failed = CampaignRecipient.query.filter_by(campaign_id=campaign_id, status="failed").all()

    retry_targets = pending + failed
    if not retry_targets:
        flash("No pending or failed recipients to retry.", "error")
        return redirect(url_for("campaigns.campaign_detail", campaign_id=campaign.id))

    quota = _daily_quota_status(current_user.id)
    if len(retry_targets) > quota["remaining"]:
        flash(
            f"Only {quota['remaining']} of {DAILY_SEND_LIMIT} daily sends remain. "
            "Open Resend Failed and schedule the remaining recipients for tomorrow.",
            "error",
        )
        return redirect(url_for("campaigns.campaign_detail", campaign_id=campaign.id))

    for r in failed:
        r.status = "pending"
        r.error  = None

    sent_ok = CampaignRecipient.query.filter_by(campaign_id=campaign_id, status="sent").count()
    campaign.sent_fail = 0
    campaign.sent_ok   = sent_ok
    campaign.total     = sent_ok + len(retry_targets)
    campaign.status    = "queued"
    db.session.commit()

    if _scheduler and _scheduler.running:
        _scheduler.add_job(
            _send_campaign_background,
            args=[campaign_id],
            id=f"send_campaign_{campaign_id}",
            replace_existing=True,
        )
    else:
        import threading
        threading.Thread(
            target=_send_campaign_background, args=[campaign_id], daemon=True
        ).start()

    flash(
        f"Retry queued for {len(retry_targets)} recipient{'s' if len(retry_targets) != 1 else ''}. "
        "Refresh this page in a moment to see progress.",
        "success",
    )
    return redirect(url_for("campaigns.campaign_detail", campaign_id=campaign.id))

@campaigns_bp.route("/campaign/<int:campaign_id>/export.csv")
@login_required
@subscription_required
def campaign_export(campaign_id):
    campaign = Campaign.query.filter_by(id=campaign_id, user_id=current_user.id).first_or_404()
    recipients = CampaignRecipient.query.filter_by(campaign_id=campaign.id).all()
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["email", "name", "status", "sent_at", "error"])
    for r in recipients:
        writer.writerow([
            r.email,
            r.name or "",
            r.status,
            r.sent_at.isoformat() if r.sent_at else "",
            r.error or "",
        ])
    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", campaign.name or "campaign").strip("_") or "campaign"
    return Response(
        buf.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}_recipients.csv"'},
    )

@campaigns_bp.route("/extract", methods=["POST"])
@login_required
@subscription_required
def extract():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded."}), 400
    f = request.files["file"]
    filename = f.filename.lower()
    text = ""

    try:
        if filename.endswith(".xlsx") or filename.endswith(".xls"):
            wb = openpyxl.load_workbook(f, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value:
                            text += str(cell.value) + " "
        elif filename.endswith(".csv"):
            raw = f.read()
            try:
                decoded = raw.decode("utf-8")
            except UnicodeDecodeError:
                decoded = raw.decode("latin-1")
            reader = csv.reader(io.StringIO(decoded))
            for row in reader:
                text += " ".join(row) + " "
        elif filename.endswith(".pdf"):
            import pdfplumber
            with pdfplumber.open(f) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + " "
        elif filename.endswith(".docx"):
            from docx import Document
            doc = Document(f)
            for para in doc.paragraphs:
                text += para.text + " "
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
        else:
            raw = f.read()
            try:
                text = raw.decode("utf-8")
            except UnicodeDecodeError:
                text = raw.decode("latin-1")
    except Exception as e:
        return jsonify({"error": f"Could not parse file: {e}"}), 400

    emails = extract_emails(text)
    names = {e: derive_name_from_email(e) for e in emails}
    return jsonify({"emails": emails, "names": names, "total": len(emails)})

@campaigns_bp.route("/extract-url", methods=["POST"])
@login_required
@subscription_required
def extract_url():
    data = request.get_json()
    url = (data or {}).get("url", "").strip()
    if not url:
        return jsonify({"error": "URL is required."}), 400
    if not url.startswith(("http://", "https://")):
        url = "https://" + url
    
    from helpers.emails import _is_ssrf_safe
    if not _is_ssrf_safe(url):
        return jsonify({"error": "That URL is not allowed."}), 400
    try:
        from bs4 import BeautifulSoup
        resp = _safe_fetch(url)
        soup = BeautifulSoup(resp.text, "html.parser")
        text = soup.get_text(separator=" ")
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"Could not fetch URL: {e}"}), 400

    emails = extract_emails(text)
    names = {e: derive_name_from_email(e) for e in emails}
    return jsonify({"emails": emails, "names": names, "total": len(emails)})

@campaigns_bp.route("/generate", methods=["POST"])
@login_required
@subscription_required
def generate():
    OPENAI_API_KEY = os.environ.get("OPENAI_API_KEY", "")
    if not OPENAI_API_KEY:
        return jsonify({"error": "OpenAI is not configured. Contact support."}), 503

    data = request.get_json()
    brief = (data or {}).get("brief", "").strip()
    if not brief:
        return jsonify({"error": "Campaign brief is required."}), 400

    try:
        from openai import OpenAI
        client = OpenAI(api_key=OPENAI_API_KEY)
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are a marketing copywriter. Write a concise, engaging marketing email. "
                        "Return ONLY valid JSON with keys 'subject' (string, max 80 chars) and 'body' (string, plain text, 150-300 words). "
                        "No markdown, no extra keys, no explanations."
                    )
                },
                {
                    "role": "user",
                    "content": f"Campaign brief: {brief}"
                }
            ],
            response_format={"type": "json_object"},
            max_tokens=600,
        )
        result = json.loads(response.choices[0].message.content)
        if "subject" not in result or "body" not in result:
            raise ValueError("Missing keys in AI response")
        return jsonify(result)
    except Exception as e:
        return jsonify({"error": f"AI generation failed: {e}"}), 500

@campaigns_bp.route("/send-bulk", methods=["POST"])
@login_required
@subscription_required
def send_bulk():
    from extensions import _pending_attachments, _scheduler
    from helpers.scheduler_tasks import _send_campaign_background

    emails_raw    = request.form.get("emails", "[]")
    names_raw     = request.form.get("names", "{}")
    subject       = request.form.get("subject", "").strip()
    body          = request.form.get("body", "").strip()
    campaign_name = (request.form.get("name", "Campaign") or "Campaign").strip()

    try:
        emails = json.loads(emails_raw)
    except Exception:
        return jsonify({"error": "Invalid emails payload."}), 400
    if not isinstance(emails, list):
        return jsonify({"error": "Invalid emails payload."}), 400
    emails = [str(email).strip().lower() for email in emails if str(email).strip()]
    emails = list(dict.fromkeys(emails))

    try:
        names_map = json.loads(names_raw)
        if not isinstance(names_map, dict):
            names_map = {}
    except Exception:
        names_map = {}

    body_text = re.sub(r"<[^>]+>", "", body).strip()
    if not emails or not subject or not body_text:
        return jsonify({"error": "emails, subject, and body are required."}), 400

    if not current_user.smtp_host:
        return jsonify({"error": "SMTP is not configured. Go to Settings."}), 400

    quota = _daily_quota_status(current_user.id)
    schedule_overflow = request.form.get("schedule_overflow") == "true"
    overflow_count = max(len(emails) - quota["remaining"], 0)
    if overflow_count and not schedule_overflow:
        return jsonify({
            "code": "daily_limit_exceeded",
            "error": (
                f"RushMail allows {DAILY_SEND_LIMIT} email attempts per UTC day. "
                f"You have {quota['remaining']} remaining today. "
                f"Schedule the other {overflow_count} for tomorrow."
            ),
            "daily_limit": DAILY_SEND_LIMIT,
            "used_today": quota["used"],
            "remaining": quota["remaining"],
            "overflow": overflow_count,
            "scheduled_for": _utc_tomorrow_start().isoformat(),
        }), 429

    send_emails = emails[:quota["remaining"]]
    overflow_emails = emails[quota["remaining"]:]
    overflow_schedule = None
    if overflow_emails:
        overflow_schedule = _create_quota_overflow_schedule(
            current_user.id,
            campaign_name,
            subject,
            body,
            overflow_emails,
            names_map,
        )

    if not send_emails:
        db.session.commit()
        return jsonify({
            "ok": True,
            "scheduled_only": True,
            "scheduled_count": len(overflow_emails),
            "scheduled_for": overflow_schedule.next_run_at.isoformat(),
            "message": (
                f"Today's {DAILY_SEND_LIMIT}-email limit is reached. "
                f"Scheduled all {len(overflow_emails)} recipients for tomorrow."
            ),
        })

    attachments = []
    for att_file in request.files.getlist("attachments"):
        if att_file and att_file.filename:
            attachments.append((att_file.filename, att_file.read(),
                                att_file.mimetype or "application/octet-stream"))

    campaign = Campaign(
        user_id=current_user.id,
        name=campaign_name,
        subject=subject,
        body=body,
        total=len(send_emails),
        status="queued",
    )
    db.session.add(campaign)
    db.session.flush()

    for email in send_emails:
        db.session.add(CampaignRecipient(
            campaign_id=campaign.id,
            email=email,
            name=resolve_recipient_name(email, names_map) or None,
        ))

    db.session.commit()
    campaign_id = campaign.id

    if attachments:
        _pending_attachments[campaign_id] = attachments

    if _scheduler and _scheduler.running:
        _scheduler.add_job(
            _send_campaign_background,
            args=[campaign_id],
            id=f"send_campaign_{campaign_id}",
            replace_existing=True,
        )
    else:
        import threading
        threading.Thread(
            target=_send_campaign_background, args=[campaign_id], daemon=True
        ).start()

    return jsonify({
        "campaign_id": campaign_id,
        "total": len(send_emails),
        "scheduled_count": len(overflow_emails),
        "scheduled_for": overflow_schedule.next_run_at.isoformat() if overflow_schedule else None,
    })

@campaigns_bp.route("/campaign/<int:campaign_id>/status")
@login_required
def campaign_status(campaign_id):
    campaign = Campaign.query.filter_by(
        id=campaign_id, user_id=current_user.id
    ).first_or_404()
    sent    = CampaignRecipient.query.filter_by(campaign_id=campaign_id, status="sent").count()
    failed  = CampaignRecipient.query.filter_by(campaign_id=campaign_id, status="failed").count()
    pending = CampaignRecipient.query.filter_by(campaign_id=campaign_id, status="pending").count()
    if sent + failed + pending == 0 and campaign.total:
        sent = campaign.sent_ok or 0
        failed = campaign.sent_fail or 0
        pending = max(campaign.total - sent - failed, 0)

    payload = {
        "campaign_id": campaign_id,
        "status":      campaign.status,
        "total":       max(campaign.total or 0, sent + failed + pending),
        "sent":        sent,
        "failed":      failed,
        "pending":     pending,
    }
    attempted = sent + failed
    payload["sent_rate"] = round(sent / attempted * 100, 1) if attempted else None
    payload["failure_rate"] = round(failed / attempted * 100, 1) if attempted else None
    payload["can_retry"] = (
        failed + pending > 0 and campaign.status not in ("queued", "sending")
    )

    if campaign.status == "completed":
        all_recipients = CampaignRecipient.query.filter_by(campaign_id=campaign_id).all()
        payload["recipients"] = [
            {
                "email":  r.email,
                "name":   r.name or "",
                "status": r.status,
                "error":  r.error or "",
            }
            for r in all_recipients
        ]

    return jsonify(payload)
